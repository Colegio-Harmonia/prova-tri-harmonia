from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Guia_ativacao_notificacoes_Google_Chat.docx"
LOGO = ROOT / "public" / "brand" / "harmonia-logo-color.png"
GREEN = "008649"
INK = "163126"
MUTED = "52635B"
PALE_GREEN = "EAF6EF"
PALE_GRAY = "F4F7F5"


def shade(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            set_cell_margins(cell)


def add_run(paragraph, text, *, bold=False, color=INK, size=10.5):
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def format_paragraph(paragraph, after=6, before=0):
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.line_spacing = 1.25


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph()
    format_paragraph(paragraph, after=7 if level == 2 else 10, before=14 if level == 2 else 18)
    add_run(paragraph, text, bold=True, color=GREEN if level == 1 else INK, size=16 if level == 1 else 13)
    return paragraph


def add_body(doc, text, *, bold_prefix=None):
    paragraph = doc.add_paragraph()
    format_paragraph(paragraph)
    if bold_prefix:
        add_run(paragraph, bold_prefix + " ", bold=True)
        add_run(paragraph, text)
    else:
        add_run(paragraph, text)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    format_paragraph(paragraph, after=4)
    add_run(paragraph, text)
    return paragraph


def add_note(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.cell(0, 0)
    shade(cell, PALE_GREEN)
    set_cell_margins(cell, top=150, start=180, bottom=150, end=180)
    paragraph = cell.paragraphs[0]
    format_paragraph(paragraph, after=0)
    add_run(paragraph, title + " ", bold=True, color=GREEN)
    add_run(paragraph, text)


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("COLÉGIO HARMONIA  |  PROVA-TRI")
    header_run.font.name = "Calibri"
    header_run.font.size = Pt(8.5)
    header_run.font.color.rgb = RGBColor.from_string(MUTED)

    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(5)
        p.add_run().add_picture(str(LOGO), width=Inches(1.05))

    title = doc.add_paragraph()
    format_paragraph(title, after=3)
    add_run(title, "Ative suas notificações no Google Chat", bold=True, color=INK, size=20)
    subtitle = doc.add_paragraph()
    format_paragraph(subtitle, after=12)
    add_run(subtitle, "Guia rápido para professores do Prova-tri", color=MUTED, size=11)

    add_note(doc, "IMPORTANTE:", "As notificações são privadas. Você receberá apenas avisos das provas atribuídas a você.")

    add_heading(doc, "Ative em cinco passos")
    steps = [
        ("1", "Abra o Google Chat com seu e-mail institucional."),
        ("2", "No menu Apps, abra Prova-tri."),
        ("3", "Envie a mensagem Teste para o app."),
        ("4", "Clique em Conectar minha conta."),
        ("5", "Entre no Prova-tri com sua conta institucional e aguarde a confirmação."),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    set_table_widths(table, [0.48, 5.86])
    for number, description in steps:
        row = table.add_row()
        number_cell, detail_cell = row.cells
        shade(number_cell, GREEN)
        number_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        num = number_cell.paragraphs[0]
        num.alignment = WD_ALIGN_PARAGRAPH.CENTER
        format_paragraph(num, after=0)
        add_run(num, number, bold=True, color="FFFFFF", size=12)
        shade(detail_cell, PALE_GRAY)
        detail = detail_cell.paragraphs[0]
        format_paragraph(detail, after=0)
        add_run(detail, description)

    add_body(doc, "Depois de ativado, uma nova mensagem ao app deve responder: ", bold_prefix="Confirmação esperada:")
    quote = doc.add_paragraph()
    quote.paragraph_format.left_indent = Inches(0.22)
    quote.paragraph_format.space_after = Pt(8)
    add_run(quote, '“Suas notificações privadas do Prova-tri já estão ativas.”', bold=True, color=GREEN)

    add_heading(doc, "O que você receberá", level=2)
    add_bullet(doc, "Aviso de prova atribuída para revisão.")
    add_bullet(doc, "Link direto para abrir a prova no Prova-tri.")
    add_bullet(doc, "Somente mensagens relacionadas às suas atribuições.")

    add_heading(doc, "Se algo não funcionar", level=2)
    add_bullet(doc, "Se não aparecer o cartão de conexão, envie Teste novamente ao app.")
    add_bullet(doc, "O endereço correto é prova.colegioharmonia.com.br. Se abrir localhost, feche a aba e envie Teste novamente.")
    add_bullet(doc, "Se continuar sem confirmação, envie uma captura de tela para a coordenação.")

    add_heading(doc, "Teste com outro professor", level=2)
    add_body(doc, "A coordenação pode testar com outro professor agora:")
    test_steps = [
        "O professor segue os cinco passos de ativação.",
        "A coordenação atribui uma prova de teste apenas a ele.",
        "O professor confirma o recebimento da mensagem privada e abre o link.",
    ]
    for item in test_steps:
        add_bullet(doc, item)
    add_note(doc, "NÃO USE GRUPOS:", "Não é necessário criar grupo ou espaço coletivo no Google Chat.")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Prova-tri | Notificações privadas no Google Chat")
    footer_run.font.name = "Calibri"
    footer_run.font.size = Pt(8.5)
    footer_run.font.color.rgb = RGBColor.from_string(MUTED)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
